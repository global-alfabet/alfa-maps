"""
Turn a surveyed general-plan PDF + a stands spreadsheet into stand polygons in metres.

Pipeline
1. Read the spreadsheet: stand number, size (m2) and optional price / status columns.
2. Read every PDF page's text and rebuild stand-number labels, including rotated ones.
3. Pick the page that is the actual plan (most labels that sit inside closed parcels).
4. Rasterise that page, split the white space into connected regions (parcels, roads,
   open land) and match each stand label to the parcel that contains it.
5. Calibrate the drawing scale from the spreadsheet sizes (median of area_px / area_m2).
6. Classify every boundary of every stand: road, neighbouring stand or open land.
"""
from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

import cv2
import numpy as np
import pdfplumber
import pypdfium2 as pdfium

Progress = Callable[[str, float], None]

TARGET_LONG_SIDE_PX = 4800          # raster resolution for the plan page
WHITE_THRESHOLD = 200               # grey level above which a pixel counts as paper


# ----------------------------------------------------------------------------- spreadsheet
HEADER_PATTERNS = {
    "stand": re.compile(r"stand|plot|erf|lot|parcel|number|no\.?$", re.I),
    "size": re.compile(r"size|area|m2|m²|sqm|sq\.?\s*m|extent", re.I),
    "price": re.compile(r"price|amount|cost|value|usd|\$", re.I),
    "status": re.compile(r"status|state|availability|sold", re.I),
}


def _num(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("m²", "").replace("m2", "").replace("$", "")
    try:
        return float(s)
    except ValueError:
        return None


def _rows_from_file(path: Path) -> list[list]:
    if path.suffix.lower() in (".csv", ".txt"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [row for row in csv.reader(f)]
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
        rows.append([])  # sheet separator
    return rows


def read_stand_sheet(path: Path) -> list[dict]:
    """Return [{'no': '34603', 'size': 673.0, 'price': None, 'status': None}, ...]."""
    rows = _rows_from_file(path)
    out: dict[str, dict] = {}
    cols = None
    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        texts = [str(c).strip() if c is not None else "" for c in row]
        # header row?  (mostly text, contains a stand-ish and a size-ish label)
        if sum(1 for c in row if _num(c) is None and c not in (None, "")) >= 2:
            found = {}
            for i, t in enumerate(texts):
                for key, pat in HEADER_PATTERNS.items():
                    if key not in found and t and pat.search(t) and not (key == "stand" and t == "#"):
                        found[key] = i
                        break
            if "stand" in found and "size" in found:
                cols = found
            continue
        if cols:
            no = _num(row[cols["stand"]]) if cols["stand"] < len(row) else None
            size = _num(row[cols["size"]]) if cols["size"] < len(row) else None
            if no is None or size is None or no <= 0 or not (20 <= size <= 200000):
                continue
            rec = {"no": str(int(no)) if float(no).is_integer() else str(no), "size": size, "price": None, "status": None}
            if "price" in cols and cols["price"] < len(row):
                rec["price"] = _num(row[cols["price"]])
            if "status" in cols and cols["status"] < len(row) and row[cols["status"]]:
                rec["status"] = str(row[cols["status"]]).strip().lower()
            out.setdefault(rec["no"], rec)
        else:
            # no header yet: take the first integer >= 100 as stand, next plausible number as size
            nums = [(i, _num(c)) for i, c in enumerate(row) if _num(c) is not None]
            stand = next(((i, v) for i, v in nums if v >= 100 and float(v).is_integer()), None)
            if not stand:
                continue
            size = next((v for i, v in nums if i > stand[0] and 20 <= v <= 200000), None)
            if size:
                no = str(int(stand[1]))
                out.setdefault(no, {"no": no, "size": size, "price": None, "status": None})
    return list(out.values())


# ----------------------------------------------------------------------------- PDF labels
@dataclass
class Label:
    text: str
    x: float      # PDF points, top-left origin (pdfplumber)
    y: float
    size: float


def _group_chars(chars: list[dict]) -> list[Label]:
    """Rebuild strings from characters in content-stream order. Works for rotated text."""
    groups, cur = [], []

    def centre(c):
        return ((c["x0"] + c["x1"]) / 2, (c["top"] + c["bottom"]) / 2)

    for c in chars:
        if cur:
            a, b = centre(cur[-1]), centre(c)
            if math.dist(a, b) < 2.2 * max(c["size"], 0.5) and abs(c["size"] - cur[-1]["size"]) < 1e-3:
                cur.append(c)
                continue
            groups.append(cur)
        cur = [c]
    if cur:
        groups.append(cur)
    out = []
    for g in groups:
        t = "".join(ch["text"] for ch in g).strip()
        if not t:
            continue
        xs = [centre(ch)[0] for ch in g]
        ys = [centre(ch)[1] for ch in g]
        out.append(Label(t, sum(xs) / len(xs), sum(ys) / len(ys), g[0]["size"]))
    return out


def page_labels(pdf_path: Path) -> list[tuple[int, float, float, list[Label]]]:
    """[(page_index, width_pt, height_pt, labels)] for each page."""
    res = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for i, p in enumerate(pdf.pages):
            res.append((i, float(p.width), float(p.height), _group_chars(p.chars)))
    return res


def find_stand_labels(labels: list[Label], wanted: set[str]) -> dict[str, list[Label]]:
    hits: dict[str, list[Label]] = {}
    for lb in labels:
        if lb.text in wanted:
            hits.setdefault(lb.text, []).append(lb)
    # second pass: numbers split into two text runs ("345" + "66")
    missing = wanted - hits.keys()
    if missing:
        for a, b in zip(labels, labels[1:]):
            joined = a.text + b.text
            if joined in missing and math.dist((a.x, a.y), (b.x, b.y)) < 6 * max(a.size, 1):
                hits.setdefault(joined, []).append(Label(joined, (a.x + b.x) / 2, (a.y + b.y) / 2, a.size))
    return hits


# ----------------------------------------------------------------------------- raster analysis
def render_page(pdf_path: Path, page_index: int) -> tuple[np.ndarray, float]:
    doc = pdfium.PdfDocument(str(pdf_path))
    page = doc[page_index]
    w, h = page.get_size()
    scale = TARGET_LONG_SIDE_PX / max(w, h)
    img = page.render(scale=scale, grayscale=True).to_numpy()
    if img.ndim == 3:
        img = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    return img, scale


class PlanRaster:
    def __init__(self, gray: np.ndarray, scale: float):
        self.gray = gray
        self.scale = scale  # px per PDF point
        white = (gray > WHITE_THRESHOLD).astype(np.uint8)
        self.n, self.lab, self.stats, _ = cv2.connectedComponentsWithStats(white, connectivity=4)
        self.area = self.stats[:, cv2.CC_STAT_AREA]
        self.area[0] = 0  # background label 0 = lines/ink
        self.min_region = max(400, int((scale * 72 / 400) ** 2 * 1500))

    def region_at_label(self, lb: Label) -> int:
        """Component id of the parcel containing a label (searches around the text)."""
        cx, cy = int(lb.x * self.scale), int(lb.y * self.scale)
        H, W = self.lab.shape
        best = 0
        for r in range(0, int(18 * self.scale / 5.5) + 60, 2):
            for dx, dy in ((r, 0), (-r, 0), (0, r), (0, -r), (r, r), (-r, -r), (r, -r), (-r, r)):
                x, y = cx + dx, cy + dy
                if 0 <= x < W and 0 <= y < H:
                    c = self.lab[y, x]
                    if c and self.area[c] >= self.min_region:
                        return int(c)
        return best

    def region_polygon(self, comp: int, px_per_m: Optional[float] = None) -> np.ndarray:
        ppm = px_per_m or self.scale * 72 / 400 * 2.7
        x, y, w, h = self.stats[comp, :4]
        pad = int(ppm * 3) + 4
        x0, y0 = max(0, x - pad), max(0, y - pad)
        x1, y1 = min(self.lab.shape[1], x + w + pad), min(self.lab.shape[0], y + h + pad)
        m = (self.lab[y0:y1, x0:x1] == comp).astype(np.uint8) * 255
        # fill text holes, then close notches where lettering touches the boundary line
        cnts, _ = cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_NONE)
        if not cnts:
            return np.zeros((0, 2))
        m = np.zeros_like(m)
        cv2.drawContours(m, [max(cnts, key=cv2.contourArea)], -1, 255, -1)
        k = int(ppm * 2.5) | 1
        m = cv2.morphologyEx(m, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (k, k)))
        # grow by half a line width so the outline sits on the surveyed boundary line
        m = cv2.dilate(m, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (3, 3)))
        cnts, _ = cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_NONE)
        c = max(cnts, key=cv2.contourArea)
        ap = cv2.approxPolyDP(c, ppm * 0.9, True)[:, 0, :].astype(float)
        ap[:, 0] += x0
        ap[:, 1] += y0
        return _merge_short(ap, ppm * 1.2)


def _merge_short(p: np.ndarray, min_len: float) -> np.ndarray:
    """Drop vertices that create very short edges (keeps corners that matter)."""
    pts = [q for q in p]
    changed = True
    while changed and len(pts) > 3:
        changed = False
        for i in range(len(pts)):
            a, b = pts[i], pts[(i + 1) % len(pts)]
            if np.hypot(*(b - a)) < min_len:
                # remove whichever endpoint deviates least from a straight line
                prev, nxt = pts[i - 1], pts[(i + 2) % len(pts)]
                da = _dev(prev, a, b)
                db = _dev(a, b, nxt)
                del pts[(i + 1) % len(pts) if db < da else i]
                changed = True
                break
    return np.array(pts)


def _dev(p, q, r) -> float:
    d = r - p
    L = np.hypot(*d) or 1
    return abs(d[0] * (q[1] - p[1]) - d[1] * (q[0] - p[0])) / L


def poly_area(p: np.ndarray) -> float:
    if len(p) < 3:
        return 0.0
    x, y = p[:, 0], p[:, 1]
    return 0.5 * abs(np.dot(x, np.roll(y, -1)) - np.dot(y, np.roll(x, -1)))


def _point_in_poly(pt, poly) -> bool:
    return cv2.pointPolygonTest(poly.astype(np.float32).reshape(-1, 1, 2), (float(pt[0]), float(pt[1])), False) > 0


# ----------------------------------------------------------------------------- main entry
def process(pdf_paths: list[Path], sheet_path: Path, out_dir: Path, progress: Progress = lambda m, f: None) -> dict:
    progress("Reading the stands spreadsheet", 0.05)
    sheet = read_stand_sheet(sheet_path)
    if not sheet:
        raise ValueError("No stand numbers and sizes were found in the spreadsheet. It needs a stand number column and a size column.")
    wanted = {r["no"] for r in sheet}
    size_of = {r["no"]: r["size"] for r in sheet}

    progress("Reading text on every PDF page", 0.12)
    candidates = []
    for pdf in pdf_paths:
        for idx, w, h, labels in page_labels(pdf):
            hits = find_stand_labels(labels, wanted)
            candidates.append((len(hits), pdf, idx, labels, hits))
    candidates.sort(key=lambda c: -c[0])
    if not candidates or candidates[0][0] == 0:
        raise ValueError("None of the stand numbers from the spreadsheet were found as text in the PDF. The plan must be a vector PDF (exported from CAD), not a scan.")

    # try the best few pages; keep the one whose parcel areas agree best with the sheet
    best = None
    for n_hits, pdf, idx, labels, hits in candidates[:3]:
        if n_hits < max(3, candidates[0][0] * 0.3):
            continue
        progress(f"Analysing page {idx + 1} of {pdf.name}", 0.25)
        gray, scale = render_page(pdf, idx)
        R = PlanRaster(gray, scale)
        ratios = []
        for no, lbs in hits.items():
            comp = R.region_at_label(lbs[0])
            if comp:
                ratios.append(R.area[comp] / size_of[no])
        if len(ratios) < 3:
            continue
        ratios = np.array(ratios)
        med = float(np.median(ratios))
        agree = float(np.mean(np.abs(ratios / med - 1) < 0.15))
        score = agree * len(ratios)
        if best is None or score > best[0]:
            best = (score, pdf, idx, labels, hits, R, med)
    if best is None:
        raise ValueError("Stand numbers were found, but not inside closed parcels. Check that the PDF is the surveyed layout plan.")
    _, pdf, page_idx, labels, hits, R, med = best

    progress("Matching stands to parcels", 0.45)
    px_per_m = math.sqrt(med)
    stand_comp: dict[str, int] = {}
    for no, lbs in hits.items():
        best_c, best_err = 0, 9e9
        for lb in lbs:
            c = R.region_at_label(lb)
            if not c:
                continue
            err = abs(R.area[c] / size_of[no] / med - 1)
            if err < best_err:
                best_c, best_err = c, err
        if best_c:
            stand_comp[no] = best_c
    # refine the scale using only good matches
    good = [R.area[c] / size_of[n] for n, c in stand_comp.items() if abs(R.area[c] / size_of[n] / med - 1) < 0.1]
    if len(good) >= 3:
        med = float(np.median(good))
        px_per_m = math.sqrt(med)

    progress("Tracing every parcel on the plan", 0.6)
    typical = float(np.median([R.area[c] for c in stand_comp.values()]))
    road_comp = int(np.argmax(R.area))
    comp_label: dict[int, str] = {}
    for lb in labels:
        if re.fullmatch(r"\d{3,6}[A-Z]?", lb.text):
            x, y = int(lb.x * R.scale), int(lb.y * R.scale)
            if 0 <= y < R.lab.shape[0] and 0 <= x < R.lab.shape[1]:
                c = R.lab[y, x]
                if not c:
                    c = R.region_at_label(lb)
                if c and c not in comp_label:
                    comp_label[int(c)] = lb.text
    stand_comps = set(stand_comp.values())
    parcel_comps, open_comps = [], []
    for c, text in comp_label.items():
        if c in stand_comps or c == road_comp:
            continue
        a = R.area[c]
        if 0.3 * typical <= a <= 4 * typical:
            parcel_comps.append(c)
        elif a > 4 * typical and not _is_strip(R, c, px_per_m):
            open_comps.append(c)

    # parcels whose number could not be read: same size range, inside the estate footprint
    known = stand_comps | set(parcel_comps) | set(open_comps)
    pts = np.array([[R.stats[c, 0] + R.stats[c, 2] / 2, R.stats[c, 1] + R.stats[c, 3] / 2] for c in known], dtype=np.float32)
    hull = cv2.convexHull(pts)
    for c in range(1, R.n):
        if c in known or c == road_comp:
            continue
        a = R.area[c]
        if not (0.35 * typical <= a <= 3 * typical):
            continue
        x, y, w, h = R.stats[c, :4]
        if a / float(w * h) < 0.35:
            continue
        if cv2.pointPolygonTest(hull, (float(x + w / 2), float(y + h / 2)), False) < 0:
            continue
        parcel_comps.append(c)

    # coordinates: metres, origin at centre of all traced parcels, x right, z down (image axes)
    all_c = list(stand_comps) + parcel_comps + open_comps
    xs = np.concatenate([[R.stats[c, 0], R.stats[c, 0] + R.stats[c, 2]] for c in all_c])
    ys = np.concatenate([[R.stats[c, 1], R.stats[c, 1] + R.stats[c, 3]] for c in all_c])
    ox, oy = float(xs.mean()), float(ys.mean())

    def to_m(p: np.ndarray) -> list[list[float]]:
        return [[round((x - ox) / px_per_m, 2), round((y - oy) / px_per_m, 2)] for x, y in p]

    polys_px: dict[int, np.ndarray] = {}
    for c in all_c:
        polys_px[c] = R.region_polygon(c, px_per_m)
    # final scale: calibrate on the traced outlines themselves (they sit on the boundary lines)
    r2 = np.array([poly_area(polys_px[c]) / size_of[n] for n, c in stand_comp.items() if len(polys_px[c]) >= 3])
    if len(r2) >= 3:
        m2 = float(np.median(r2))
        m2 = float(np.median(r2[np.abs(r2 / m2 - 1) < 0.1])) if np.any(np.abs(r2 / m2 - 1) < 0.1) else m2
        px_per_m = math.sqrt(m2)

    progress("Classifying boundaries (road, neighbour, open land)", 0.75)

    ink = R.gray <= WHITE_THRESHOLD

    def classify_point(x: float, y: float, dx: float, dy: float, own: int) -> tuple[str, str, float]:
        """Walk outward from a boundary. Returns (kind, label, width_m)."""
        H, W = R.lab.shape
        step = 0.25 * px_per_m
        d = 0.0
        comp, start = 0, None
        while d < 45 * px_per_m:
            d += step
            px, py = int(x + dx * d), int(y + dy * d)
            if not (0 <= px < W and 0 <= py < H):
                return "open", "", 99.0
            c = int(R.lab[py, px])
            if start is None:
                if c and c != own and R.area[c] >= R.min_region:
                    comp, start = c, d
                continue
            if ink[py, px]:
                break
        if start is None:
            return "open", "", 99.0
        width = (d - start) / px_per_m
        if comp in stand_comps:
            return "neighbour", comp_no[comp], width
        a = R.area[comp]
        if comp in comp_label and a <= 4 * typical:
            return "neighbour", comp_label[comp], width
        if a <= 4 * typical:
            return "neighbour", "", width
        if width <= 24:
            return "road", "", width
        return "open", comp_label.get(comp, "") if comp not in road_like else "", width

    comp_no = {c: n for n, c in stand_comp.items()}
    road_like = {road_comp}

    stands = []
    for rec in sheet:
        no = rec["no"]
        base = {"no": no, "size": rec["size"], "price": rec["price"], "status": _norm_status(rec["status"])}
        c = stand_comp.get(no)
        if not c:
            stands.append({**base, "quality": "missing", "poly": [], "edges": [], "measured": None})
            continue
        p = polys_px[c]
        measured = poly_area(p) / (px_per_m ** 2)
        err = abs(measured / rec["size"] - 1)
        quality = "ok" if err < 0.08 else ("warn" if err < 0.2 else "check")
        edges = []
        n = len(p)
        for i in range(n):
            a, b = p[i], p[(i + 1) % n]
            L = float(np.hypot(*(b - a)))
            dx, dy = (b - a) / (L or 1)
            nx, ny = dy, -dx
            mid = (a + b) / 2
            if _point_in_poly(mid + np.array([nx, ny]) * 2, p):
                nx, ny = -nx, -ny
            votes: dict[tuple[str, str], int] = {}
            widths = []
            for t in (0.2, 0.35, 0.5, 0.65, 0.8):
                q = a + (b - a) * t
                kind_, label_, w_ = classify_point(q[0], q[1], nx, ny, c)
                votes[(kind_, label_)] = votes.get((kind_, label_), 0) + 1
                if kind_ == "road" and w_ > 5:
                    widths.append(w_)
            kind, label = max(votes.items(), key=lambda kv: kv[1])[0]
            e = {"kind": kind, "label": label, "length": round(L / px_per_m, 2)}
            if kind == "road" and widths:
                e["width"] = round(float(np.median(widths)))
            edges.append(e)
        stands.append({**base, "quality": quality, "measured": round(measured, 1), "poly": to_m(p), "edges": edges})

    parcels = [{"label": comp_label.get(c, ""), "poly": to_m(polys_px[c])} for c in parcel_comps if len(polys_px[c]) >= 3]
    opens = [{"label": comp_label.get(c, ""), "poly": to_m(polys_px[c])} for c in open_comps if len(polys_px[c]) >= 3]

    progress("Saving the plan image", 0.9)
    H, W = R.gray.shape
    k = min(1.0, 3000 / max(H, W))
    small = cv2.resize(R.gray, (int(W * k), int(H * k)), interpolation=cv2.INTER_AREA)
    cv2.imwrite(str(out_dir / "plan.png"), small)
    # transform from metres to plan.png pixels:  px = (m * px_per_m + origin) * k
    found = sum(1 for s in stands if s["quality"] != "missing")
    return {
        "source": {"pdf": pdf.name, "page": page_idx + 1},
        "scale": {"px_per_m": px_per_m, "origin": [ox, oy], "image_k": k, "image_size": [int(W * k), int(H * k)]},
        "stands": stands,
        "parcels": parcels,
        "open": opens,
        "summary": {
            "in_sheet": len(sheet),
            "found": found,
            "ok": sum(1 for s in stands if s["quality"] == "ok"),
            "warn": sum(1 for s in stands if s["quality"] == "warn"),
            "check": sum(1 for s in stands if s["quality"] == "check"),
            "missing": sum(1 for s in stands if s["quality"] == "missing"),
            "parcels": len(parcels),
            "open": len(opens),
        },
    }


def _is_strip(R: "PlanRaster", comp: int, px_per_m: float) -> bool:
    """True for road-like regions: long and narrow (mean width = 2*area/perimeter)."""
    x, y, w, h = R.stats[comp, :4]
    m = (R.lab[y:y + h, x:x + w] == comp).astype(np.uint8)
    cnts, _ = cv2.findContours(m, cv2.RETR_CCOMP, cv2.CHAIN_APPROX_SIMPLE)
    perim = sum(cv2.arcLength(c, True) for c in cnts) or 1
    return (2 * R.area[comp] / perim) / px_per_m < 18


def _norm_status(s: Optional[str]) -> str:
    if not s:
        return "available"
    s = s.lower()
    if "sold" in s:
        return "sold"
    if "reserv" in s or "hold" in s or "pending" in s:
        return "reserved"
    return "available"
